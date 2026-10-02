# Genera presupuesto_hidraulico.sql a partir del Excel del presupuesto (hoja PRES).
# Uso: python3 generar_sql.py <archivo.xlsx> [salida.sql]   (requiere: pip install openpyxl)

import openpyxl, re, collections, unicodedata

import sys
SRC = sys.argv[1] if len(sys.argv) > 1 else 'Presupuesto_Hidraulico_Opera_Tower_v4.xlsx'
OUT = sys.argv[2] if len(sys.argv) > 2 else 'presupuesto_hidraulico.sql'

wb = openpyxl.load_workbook(SRC, data_only=False)
ws = wb['PRES']

# ---------- normalización de descripciones ----------
FIXES = [
    ('instalaicon', 'instalación'), ('instalacioón', 'instalación'),
    ('instalacion ', 'instalación '), ('Instalación', 'instalación'),
    ('Regitro', 'Registro'), ('Hchura', 'Hechura'), ('Parquoe', 'Parqueo'),
    ('transpaente', 'transparente'), ('Dimeniones', 'Dimensiones'),
    ('Temperatua', 'Temperatura'), ('Profunidad', 'Profundidad'),
    ('minino', 'mínimo'), ('ASTM D2242', 'ASTM D2241'), ('ASTM A-54', 'ASTM A-53'),
    ('ø3/4"incluye', 'ø3/4" incluye'), ('ø1"incluye', 'ø1" incluye'),
    ('ø 4"', 'ø4"'), ('1 1 /2"', '1 1/2"'), ('Tuberia', 'Tubería'), ('tuberia', 'tubería'),
    ('Valvula', 'Válvula'), ('valvula', 'válvula'), ('presion', 'presión'),
    ('Presion', 'Presión'), ('soporteria', 'soportería'), ('señaletica', 'señalética'),
    ('Cedula', 'Cédula'), ('Diametro', 'Diámetro'), ('Modulo', 'Módulo'),
    ('regulación de presión para 4"', 'regulación de presión para ø4"'),
    ('Soporteria', 'Soportería'), ('Tapon', 'Tapón'), ('Caja Septica', 'Caja Séptica'),
    ('Retencion', 'Retención'), ('Linea', 'Línea'), ('Succion', 'Succión'), ('automaticos', 'automáticos'),
    ('Plasticos', 'Plásticos'), ('Instantáneo', 'instantáneo'), ('Quimico', 'Químico'),
]
def norm(s):
    s = s.strip()
    s = re.sub(r'\s+', ' ', s)
    s = s.rstrip('.,').strip()
    for a, b in FIXES:
        s = s.replace(a, b)
    s = s.replace(' ,', ',').replace(' .', '.')
    return s

def nounit(u):
    u = (u or '').strip().upper()
    return 'ml' if u == 'ML' else u

# ---------- lectura ----------
capitulos = []   # dict(letra, nombre, partidas=[...])
cap = part = None
for r in range(5, ws.max_row + 1):
    d, e = ws.cell(r, 4).value, ws.cell(r, 5).value
    if e is None:
        continue
    if isinstance(d, str) and not d.startswith('='):
        cap = dict(letra=d.strip(), nombre=norm(e), partidas=[])
        capitulos.append(cap)
    elif isinstance(d, int):
        part = dict(numero_origen=d, nombre=norm(e), items=[])
        cap['partidas'].append(part)
    elif isinstance(d, str) and d.startswith('='):
        qty = ws.cell(r, 9).value
        part['items'].append(dict(desc=norm(e), unidad=nounit(ws.cell(r, 10).value),
                                  cantidad=float(qty) if qty is not None else 0.0, fila=r))

# ---------- catálogo de conceptos ----------
SIST = {'A': 'AP', 'B': 'AN', 'C': 'AL', 'D': 'CI'}
sistemas = [('AP', 'Red de agua potable fría y caliente'), ('AN', 'Red de aguas negras'),
            ('AL', 'Red de aguas lluvias'), ('CI', 'Sistema de combate contra incendio')]
unidades = [('ml', 'Metro lineal'), ('UNIDAD', 'Unidad'), ('SG', 'Suma global')]

def tipo(desc):
    d = desc.lower()
    tests = [
        ('montante principal', 'Montante'), ('bajada de agua potable', 'Bajada de agua potable'),
        ('kit de válvula', 'Kit de válvulas'), ('control de zona', 'Control de zona'),
        ('válvula para conexión de cuerpo de bomberos', 'Conexión de bomberos'),
        ('bombas trituradoras', 'Bomba de aguas negras'),
        ('bombas verticales', 'Equipo de bombeo agua potable'), ('bomba estacionaria para incendio', 'Equipo de bombeo contra incendio'),
        ('válvula', 'Válvula'), ('grifo', 'Grifo'), ('medidor', 'Medidor'),
        ('calentador', 'Calentador'), ('rociador', 'Rociador'), ('gabinete', 'Gabinete contra incendio'),
        ('tapón de registro', 'Tapón de registro'), ('caja septica', 'Caja séptica'), ('caja séptica', 'Caja séptica'),
        ('cajas de registro', 'Caja de registro'), ('bombas trituradoras', 'Bomba de aguas negras'),
        ('bombas verticales', 'Equipo de bombeo agua potable'), ('bomba estacionaria para incendio', 'Equipo de bombeo contra incendio'),
        ('soportería vertical', 'Soportería'), ('coladera', 'Coladera'), ('canal para aguas lluvias', 'Canal'),
        ('retencion pluvial', 'Retención pluvial'), ('retención pluvial', 'Retención pluvial'),
        ('pozo de visita', 'Pozo de visita'), ('tubería', 'Tubería'),
    ]
    for k, v in tests:
        if k in d: return v
    return 'Otro'

def material(desc):
    d = desc.lower()
    if 'novafort' in d: return 'PVC-U Novafort'
    if 'cpvc' in d: return 'CPVC'
    if 'acero al carb' in d: return 'Acero al carbón'
    if 'pvc' in d: return 'PVC'
    if 'inoxidable' in d: return 'Acero inoxidable'
    if 'bronce' in d: return 'Bronce'
    if 'lamina galvanizada' in d or 'lámina galvanizada' in d: return 'Lámina galvanizada'
    if 'concreto' in d: return 'Concreto reforzado'
    if 'block' in d: return 'Block'
    if 'ladrillo' in d: return 'Ladrillo'
    if 'plastico' in d or 'plásticos' in d: return 'Plástico'
    return None

def diametro(desc):
    m = re.search(r'ø\s*(\d+(?: \d+/\d+)?|\d+/\d+)"', desc)
    if m: return m.group(1) + '"'
    m = re.search(r'tubería (\d+)"', desc, re.I)
    if m: return m.group(1) + '"'
    return None

def norma(desc):
    m = re.search(r'ASTM [A-Z]-?\d+', desc)
    return m.group(0) if m else None

conceptos = {}   # desc -> dict
orden_por_sistema = collections.Counter()
for cap in capitulos:
    for part in cap['partidas']:
        for it in part['items']:
            key = it['desc']
            if key not in conceptos:
                sis = SIST[cap['letra']]
                if key.startswith('Bombas Verticales en L'): sis = 'AP'
                unidad = it['unidad']
                # corrección: tubería medida en UNIDAD en el origen -> ml
                if tipo(key) == 'Tubería' and unidad == 'UNIDAD':
                    unidad = 'ml'
                orden_por_sistema[sis] += 1
                conceptos[key] = dict(codigo=f'{sis}-{orden_por_sistema[sis]:03d}', sistema=sis,
                                      unidad=unidad, tipo=tipo(key), material=material(key),
                                      diametro=diametro(key), norma=norma(key), usos=0)
            conceptos[key]['usos'] += 1
            it['codigo'] = conceptos[key]['codigo']

# ---------- escritura SQL ----------
def q(s):
    if s is None: return 'NULL'
    return "'" + str(s).replace('\\', '\\\\').replace("'", "''") + "'"

L = []
w = L.append
w('-- =====================================================================')
w('-- BASE DE DATOS DE PRESUPUESTOS HIDRAULICOS')
w('-- Generado a partir de: Presupuesto_Hidraulico_Opera_Tower_v4.xlsx (hoja PRES)')
w('-- Motor objetivo: MySQL 8 / MariaDB 10 (XAMPP). Importar desde phpMyAdmin o:')
w('--   mysql -u root -p < presupuesto_hidraulico.sql')
w('-- Alcance: estructura del presupuesto, catalogo de conceptos y cantidades.')
w('-- NO incluye precios unitarios (se trabajan en un proceso aparte).')
w('-- =====================================================================')
w('')
w('CREATE DATABASE IF NOT EXISTS presupuestos_hidraulicos')
w('  CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci;')
w('USE presupuestos_hidraulicos;')
w('')
w('SET FOREIGN_KEY_CHECKS = 0;')
w('DROP TABLE IF EXISTS presupuesto_detalle;')
w('DROP TABLE IF EXISTS partidas;')
w('DROP TABLE IF EXISTS capitulos;')
w('DROP TABLE IF EXISTS proyectos;')
w('DROP TABLE IF EXISTS conceptos;')
w('DROP TABLE IF EXISTS sistemas;')
w('DROP TABLE IF EXISTS unidades;')
w('SET FOREIGN_KEY_CHECKS = 1;')
w('')
w('-- ---------------------------------------------------------------------')
w('-- 1. CATALOGOS (reutilizables entre proyectos)')
w('-- ---------------------------------------------------------------------')
w('''CREATE TABLE unidades (
  id        INT AUTO_INCREMENT PRIMARY KEY,
  codigo    VARCHAR(10)  NOT NULL UNIQUE,   -- ml, UNIDAD, SG
  nombre    VARCHAR(50)  NOT NULL
) ENGINE=InnoDB;''')
w('')
w('''CREATE TABLE sistemas (
  id        INT AUTO_INCREMENT PRIMARY KEY,
  codigo    VARCHAR(5)   NOT NULL UNIQUE,   -- AP, AN, AL, CI
  nombre    VARCHAR(100) NOT NULL
) ENGINE=InnoDB;''')
w('')
w('''-- Catalogo maestro de conceptos: una fila por descripcion de partida.
-- Las columnas tipo_elemento / material / diametro / norma se derivan de la
-- descripcion para poder filtrar y agrupar; la descripcion es el texto oficial.
CREATE TABLE conceptos (
  id             INT AUTO_INCREMENT PRIMARY KEY,
  codigo         VARCHAR(10)  NOT NULL UNIQUE,  -- ej. AP-001
  sistema_id     INT          NOT NULL,
  unidad_id      INT          NOT NULL,
  descripcion    TEXT         NOT NULL,
  tipo_elemento  VARCHAR(50)  NULL,
  material       VARCHAR(50)  NULL,
  diametro       VARCHAR(15)  NULL,
  norma          VARCHAR(20)  NULL,
  activo         TINYINT(1)   NOT NULL DEFAULT 1,
  CONSTRAINT fk_conceptos_sistema FOREIGN KEY (sistema_id) REFERENCES sistemas(id),
  CONSTRAINT fk_conceptos_unidad  FOREIGN KEY (unidad_id)  REFERENCES unidades(id)
) ENGINE=InnoDB;''')
w('')
w('-- ---------------------------------------------------------------------')
w('-- 2. ESTRUCTURA DEL PRESUPUESTO (por proyecto)')
w('-- ---------------------------------------------------------------------')
w('''CREATE TABLE proyectos (
  id           INT AUTO_INCREMENT PRIMARY KEY,
  nombre       VARCHAR(150) NOT NULL,
  version      VARCHAR(20)  NULL,
  descripcion  VARCHAR(255) NULL,
  creado_en    TIMESTAMP    NOT NULL DEFAULT CURRENT_TIMESTAMP
) ENGINE=InnoDB;''')
w('')
w('''-- Capitulo = letra del presupuesto (A, B, C, D)
CREATE TABLE capitulos (
  id           INT AUTO_INCREMENT PRIMARY KEY,
  proyecto_id  INT          NOT NULL,
  letra        VARCHAR(3)   NOT NULL,
  nombre       VARCHAR(150) NOT NULL,
  sistema_id   INT          NOT NULL,
  CONSTRAINT uq_capitulo UNIQUE (proyecto_id, letra),
  CONSTRAINT fk_capitulos_proyecto FOREIGN KEY (proyecto_id) REFERENCES proyectos(id) ON DELETE CASCADE,
  CONSTRAINT fk_capitulos_sistema  FOREIGN KEY (sistema_id)  REFERENCES sistemas(id)
) ENGINE=InnoDB;''')
w('')
w('''-- Partida = nivel o bloque dentro del capitulo (1, 2, 3 ... en la hoja)
CREATE TABLE partidas (
  id             INT AUTO_INCREMENT PRIMARY KEY,
  capitulo_id    INT          NOT NULL,
  orden          INT          NOT NULL,          -- orden dentro del capitulo (1..n)
  numero_origen  INT          NULL,              -- numero tal como aparece en la hoja
  nombre         VARCHAR(150) NOT NULL,
  npt            VARCHAR(30)  NULL,              -- nivel de piso terminado, si aplica
  CONSTRAINT uq_partida UNIQUE (capitulo_id, orden),
  CONSTRAINT fk_partidas_capitulo FOREIGN KEY (capitulo_id) REFERENCES capitulos(id) ON DELETE CASCADE
) ENGINE=InnoDB;''')
w('')
w('''-- Subpartida = concepto del catalogo con su cantidad en una partida.
-- El precio unitario NO vive aqui; cuando se defina ese proceso se agrega una
-- tabla precios_unitarios (concepto_id, fecha, precio, fuente) y se cruza por concepto_id.
CREATE TABLE presupuesto_detalle (
  id           INT            AUTO_INCREMENT PRIMARY KEY,
  partida_id   INT            NOT NULL,
  orden        INT            NOT NULL,          -- .01, .02, ... dentro de la partida
  concepto_id  INT            NOT NULL,
  cantidad     DECIMAL(12,2)  NOT NULL DEFAULT 0,
  observacion  VARCHAR(255)   NULL,
  CONSTRAINT uq_detalle UNIQUE (partida_id, orden),
  CONSTRAINT fk_detalle_partida  FOREIGN KEY (partida_id)  REFERENCES partidas(id) ON DELETE CASCADE,
  CONSTRAINT fk_detalle_concepto FOREIGN KEY (concepto_id) REFERENCES conceptos(id)
) ENGINE=InnoDB;''')
w('')
w('-- ---------------------------------------------------------------------')
w('-- 3. DATOS: catalogos')
w('-- ---------------------------------------------------------------------')
w('INSERT INTO unidades (codigo, nombre) VALUES')
w(',\n'.join(f'  ({q(c)}, {q(n)})' for c, n in unidades) + ';')
w('')
w('INSERT INTO sistemas (codigo, nombre) VALUES')
w(',\n'.join(f'  ({q(c)}, {q(n)})' for c, n in sistemas) + ';')
w('')
w(f'-- {len(conceptos)} conceptos distintos (de {sum(c["usos"] for c in conceptos.values())} subpartidas en la hoja)')
w('INSERT INTO conceptos (codigo, sistema_id, unidad_id, descripcion, tipo_elemento, material, diametro, norma) VALUES')
rows = []
ORD = {'AP':0,'AN':1,'AL':2,'CI':3}
for desc, c in sorted(conceptos.items(), key=lambda kv: (ORD[kv[1]['sistema']], kv[1]['codigo'])):
    rows.append(f"  ({q(c['codigo'])}, (SELECT id FROM sistemas WHERE codigo={q(c['sistema'])}), "
                f"(SELECT id FROM unidades WHERE codigo={q(c['unidad'])}), {q(desc)}, "
                f"{q(c['tipo'])}, {q(c['material'])}, {q(c['diametro'])}, {q(c['norma'])})")
w(',\n'.join(rows) + ';')
w('')
w('-- ---------------------------------------------------------------------')
w('-- 4. DATOS: proyecto OPERA TOWER')
w('-- ---------------------------------------------------------------------')
w("INSERT INTO proyectos (nombre, version, descripcion) VALUES")
w("  ('OPERA TOWER', 'v4', 'Presupuesto de obras hidráulicas - obras internas a torre de apartamentos');")
w('SET @proy = LAST_INSERT_ID();')
w('')
for cap in capitulos:
    w(f"-- Capitulo {cap['letra']}: {cap['nombre']}")
    w(f"INSERT INTO capitulos (proyecto_id, letra, nombre, sistema_id) VALUES")
    w(f"  (@proy, {q(cap['letra'])}, {q(cap['nombre'])}, (SELECT id FROM sistemas WHERE codigo={q(SIST[cap['letra']])}));")
    w(f"SET @cap_{cap['letra']} = LAST_INSERT_ID();")
    w(f"INSERT INTO partidas (capitulo_id, orden, numero_origen, nombre, npt) VALUES")
    prow = []
    for i, p in enumerate(cap['partidas'], 1):
        m = re.search(r'\(NPT\s*=\s*(.*?)\)', p['nombre'])
        npt = re.sub(r'\s+', '', m.group(1)) if m else None
        nombre = re.sub(r'\s*\(NPT.*?\)', '', p['nombre']).strip()
        p['orden'] = i
        prow.append(f"  (@cap_{cap['letra']}, {i}, {p['numero_origen']}, {q(nombre)}, {q(npt)})")
    w(',\n'.join(prow) + ';')
    w('')
w('-- Subpartidas (cantidad por concepto y partida)')
for cap in capitulos:
    for p in cap['partidas']:
        w(f"-- {cap['letra']}.{p['orden']} {p['nombre']}")
        w("INSERT INTO presupuesto_detalle (partida_id, orden, concepto_id, cantidad) VALUES")
        drow = []
        for j, it in enumerate(p['items'], 1):
            drow.append(f"  ((SELECT id FROM partidas WHERE capitulo_id=@cap_{cap['letra']} AND orden={p['orden']}), {j}, "
                        f"(SELECT id FROM conceptos WHERE codigo={q(it['codigo'])}), {it['cantidad']:.2f})")
        w(',\n'.join(drow) + ';')
w('')
w('-- ---------------------------------------------------------------------')
w('-- 5. VISTA DE CONSULTA: presupuesto desglosado tal como se lee en la hoja')
w('-- ---------------------------------------------------------------------')
w('''CREATE OR REPLACE VIEW v_presupuesto AS
SELECT
  pr.nombre                                   AS proyecto,
  ca.letra                                    AS capitulo,
  ca.nombre                                   AS capitulo_nombre,
  pa.orden                                    AS partida,
  pa.nombre                                   AS partida_nombre,
  pa.npt,
  CONCAT(pa.orden, '.', LPAD(de.orden, 2, '0')) AS subpartida,
  co.codigo                                   AS concepto,
  co.descripcion,
  un.codigo                                   AS unidad,
  de.cantidad
FROM presupuesto_detalle de
JOIN partidas  pa ON pa.id = de.partida_id
JOIN capitulos ca ON ca.id = pa.capitulo_id
JOIN proyectos pr ON pr.id = ca.proyecto_id
JOIN conceptos co ON co.id = de.concepto_id
JOIN unidades  un ON un.id = co.unidad_id;''')
w('')
w('-- Ejemplos de consulta:')
w("-- SELECT * FROM v_presupuesto WHERE capitulo = 'A' AND partida = 1;")
w("-- SELECT concepto, descripcion, unidad, SUM(cantidad) total FROM v_presupuesto GROUP BY concepto, descripcion, unidad ORDER BY concepto;")
w("-- SELECT * FROM conceptos WHERE tipo_elemento = 'Tubería' AND material = 'PVC';")
w('')

open(OUT, 'w', encoding='utf-8').write('\n'.join(L))
print('escrito', OUT, 'conceptos', len(conceptos), 'capitulos', len(capitulos),
      'partidas', sum(len(c['partidas']) for c in capitulos),
      'detalle', sum(len(p['items']) for c in capitulos for p in c['partidas']))
print('unidades vistas:', collections.Counter(c['unidad'] for c in conceptos.values()))
print('tipos:', collections.Counter(c['tipo'] for c in conceptos.values()))
